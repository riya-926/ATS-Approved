"""
Create a test resume.docx file for testing.
"""
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: python-docx not installed. Install with: pip install python-docx")
    sys.exit(1)
import sys

def create_test_resume():
    """Create a test resume with various sections."""
    doc = Document()
    
    # Header
    header = doc.add_heading('John Doe', 0)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Contact info
    contact = doc.add_paragraph('john.doe@email.com | (555) 123-4567 | linkedin.com/in/johndoe')
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_format = contact.runs[0].font
    contact_format.size = Pt(10)
    
    # Summary
    doc.add_heading('Professional Summary', level=1)
    summary = doc.add_paragraph(
        'Experienced software engineer with 5+ years developing scalable web applications. '
        'Proficient in Python, JavaScript, and cloud technologies. Strong background in '
        'REST API development and microservices architecture.'
    )
    
    # Experience
    doc.add_heading('Professional Experience', level=1)
    
    # Job 1
    job1_title = doc.add_paragraph()
    job1_title.add_run('Senior Software Engineer').bold = True
    job1_title.add_run(' | Tech Company Inc. | Jan 2021 - Present')
    
    job1_bullet1 = doc.add_paragraph('Developed REST APIs using Python and FastAPI', style='List Bullet')
    job1_bullet2 = doc.add_paragraph('Designed and implemented microservices architecture', style='List Bullet')
    job1_bullet3 = doc.add_paragraph('Optimized database queries reducing response time by 40%', style='List Bullet')
    
    # Job 2
    job2_title = doc.add_paragraph()
    job2_title.add_run('Software Engineer').bold = True
    job2_title.add_run(' | Startup Corp | Jun 2019 - Dec 2020')
    
    job2_bullet1 = doc.add_paragraph('Built frontend applications using React and TypeScript', style='List Bullet')
    job2_bullet2 = doc.add_paragraph('Implemented automated testing reducing bugs by 30%', style='List Bullet')
    # Note: Only 2 bullets to test consistency
    
    # Skills
    doc.add_heading('Technical Skills', level=1)
    skills = doc.add_paragraph('Python, JavaScript, TypeScript, React, FastAPI, PostgreSQL, AWS, Docker, Git')
    
    # Education
    doc.add_heading('Education', level=1)
    education = doc.add_paragraph()
    education.add_run('Bachelor of Science in Computer Science').bold = True
    education.add_run(' | State University | 2019')
    
    # Save
    output_path = 'test_resume.docx'
    doc.save(output_path)
    print(f"✅ Created test resume: {output_path}")
    return output_path

if __name__ == "__main__":
    create_test_resume()
