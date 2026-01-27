#!/usr/bin/env python3
"""
Create a sample resume DOCX file for testing.

This creates a simple resume with Experience, Skills, and Summary sections.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, str(Path(__file__).parent.parent))


def create_sample_resume(output_path: Path):
    """Create a sample resume DOCX file."""
    doc = Document()

    # Title
    title = doc.add_heading("John Doe", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Contact info
    contact = doc.add_paragraph("john.doe@email.com | (555) 123-4567 | linkedin.com/in/johndoe")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Summary section
    doc.add_heading("Summary", level=1)
    summary = doc.add_paragraph(
        "Experienced software engineer with 5+ years developing scalable web applications "
        "using Python and JavaScript. Proficient in React, Node.js, and cloud technologies."
    )

    # Experience section
    doc.add_heading("Experience", level=1)

    # Job 1
    doc.add_paragraph("Senior Software Engineer", style="Heading 2")
    doc.add_paragraph("Tech Company Inc. | 2020 - Present")
    exp1_bullet1 = doc.add_paragraph(
        "• Developed REST APIs and microservices using Python and FastAPI, serving 1M+ requests daily",
        style="List Bullet",
    )
    exp1_bullet2 = doc.add_paragraph(
        "• Built responsive front-end interfaces with React and TypeScript, improving user engagement by 30%",
        style="List Bullet",
    )
    exp1_bullet3 = doc.add_paragraph(
        "• Optimized database queries in PostgreSQL, reducing query time by 40%",
        style="List Bullet",
    )

    # Job 2
    doc.add_paragraph("Software Engineer", style="Heading 2")
    doc.add_paragraph("Startup Co. | 2018 - 2020")
    exp2_bullet1 = doc.add_paragraph(
        "• Implemented CI/CD pipelines using Jenkins and Docker, reducing deployment time by 50%",
        style="List Bullet",
    )
    exp2_bullet2 = doc.add_paragraph(
        "• Collaborated with cross-functional teams using Agile methodologies",
        style="List Bullet",
    )

    # Skills section
    doc.add_heading("Skills", level=1)
    skills = doc.add_paragraph("Python, JavaScript, TypeScript, React, Node.js, PostgreSQL, MongoDB, AWS, Docker, Kubernetes, Git, CI/CD")

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science", style="Heading 2")
    doc.add_paragraph("University Name | 2014 - 2018")

    # Save
    doc.save(output_path)
    print(f"✓ Created sample resume: {output_path}")


if __name__ == "__main__":
    output_path = Path("test_files/sample_resume.docx")
    output_path.parent.mkdir(exist_ok=True)
    create_sample_resume(output_path)
    print(f"\nYou can now test the pipeline with:")
    print(f"  python scripts/test_full_pipeline.py {output_path} test_files/sample_jd.txt")
